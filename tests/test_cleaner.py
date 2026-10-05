import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from cleaner import FIXED_ZIP_DATE, OFFICE_EXTENSIONS, PROPERTY_TYPES, iter_office_files, remove_office_metadata
from settings import desktop_directory, load_settings, restored_selection, save_settings


class CleanerTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.folder = Path(self.temporary.name)

    def make_document(self, kind):
        path = self.folder / ('sample.' + kind)
        if kind == 'docx':
            document = Document()
            document.add_paragraph('Тестовый документ')
            document.core_properties.author = 'PRIVATE AUTHOR'
            document.save(path)
        elif kind == 'pptx':
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = 'Тестовый слайд'
            slide.placeholders[1].text = 'Содержимое презентации'
            slide.notes_slide.notes_text_frame.text = 'Заметки докладчика'
            presentation.core_properties.author = 'PRIVATE AUTHOR'
            presentation.save(path)
        elif kind == 'xlsx':
            workbook = Workbook()
            workbook.active['A1'] = 'Тестовая таблица'
            workbook.active['A2'] = 7
            workbook.active['B2'] = '=A2*2'
            workbook.properties.creator = 'PRIVATE AUTHOR'
            workbook.save(path)
        return path

    def check_cleaned(self, path, before):
        with zipfile.ZipFile(path) as archive:
            self.assertIsNone(archive.testzip())
            self.assertEqual(archive.comment, b'')
            self.assertFalse(any(name.startswith('docProps/') for name in archive.namelist()))
            for info in archive.infolist():
                self.assertEqual(info.date_time, FIXED_ZIP_DATE)
                self.assertEqual(info.extra, b'')
                self.assertEqual(info.comment, b'')
                if info.filename not in ('[Content_Types].xml', '_rels/.rels'):
                    self.assertEqual(archive.read(info), before[info.filename])
            for rel in ET.fromstring(archive.read('_rels/.rels')):
                self.assertNotIn(rel.get('Type'), PROPERTY_TYPES)
            for entry in ET.fromstring(archive.read('[Content_Types].xml')):
                self.assertFalse(entry.get('PartName', '').startswith('/docProps/'))

    def test_real_word_powerpoint_and_excel_reopen_and_preserve_content(self):
        for kind in ('docx', 'pptx', 'xlsx'):
            with self.subTest(kind=kind):
                path = self.make_document(kind)
                original_bytes = path.read_bytes()
                with zipfile.ZipFile(path) as archive:
                    before = {name: archive.read(name) for name in archive.namelist()}
                result = remove_office_metadata(path)
                self.assertGreater(result.removed_items, 0)
                self.assertEqual(result.backup_path.read_bytes(), original_bytes)
                self.check_cleaned(path, before)
                if kind == 'docx':
                    self.assertEqual(Document(path).paragraphs[0].text, 'Тестовый документ')
                elif kind == 'pptx':
                    slides = Presentation(path).slides
                    self.assertEqual(slides[0].shapes.title.text, 'Тестовый слайд')
                    self.assertEqual(slides[0].notes_slide.notes_text_frame.text, 'Заметки докладчика')
                else:
                    workbook = load_workbook(path)
                    self.assertEqual(workbook.active['A1'].value, 'Тестовая таблица')
                    self.assertEqual(workbook.active['B2'].value, '=A2*2')
                    workbook.close()

    def test_all_extensions_preserve_macro_and_embedded_binary_bytes(self):
        template = self.make_document('xlsx')
        with zipfile.ZipFile(template, 'a') as archive:
            archive.writestr('xl/vbaProject.bin', b'opaque-macro-binary\x00\xff')
            archive.writestr('xl/embeddings/oleObject1.bin', b'embedded-content\x00\xfe')
        for extension in sorted(OFFICE_EXTENSIONS):
            with self.subTest(extension=extension):
                path = self.folder / ('file' + extension.upper())
                path.write_bytes(template.read_bytes())
                remove_office_metadata(path, make_backup=False)
                with zipfile.ZipFile(path) as archive:
                    self.assertEqual(archive.read('xl/vbaProject.bin'), b'opaque-macro-binary\x00\xff')
                    self.assertEqual(archive.read('xl/embeddings/oleObject1.bin'), b'embedded-content\x00\xfe')
                self.assertFalse(path.with_name(path.name + '.bak').exists())

    def test_backup_never_overwrites_previous_copy_and_repeat_is_valid(self):
        path = self.make_document('docx')
        first = remove_office_metadata(path)
        original_backup = first.backup_path.read_bytes()
        second = remove_office_metadata(path)
        self.assertEqual(second.removed_items, 0)
        self.assertNotEqual(first.backup_path, second.backup_path)
        self.assertEqual(first.backup_path.read_bytes(), original_backup)
        Document(path)

    def test_bad_zip_xml_and_legacy_formats_leave_original_unchanged(self):
        for name, content in [('bad.xlsx', b'not-a-zip'), ('old.xls', b'legacy'), ('old.ppt', b'legacy')]:
            path = self.folder / name
            path.write_bytes(content)
            with self.assertRaises(ValueError):
                remove_office_metadata(path)
            self.assertEqual(path.read_bytes(), content)
        path = self.make_document('pptx')
        broken = self.folder / 'broken.pptx'
        with zipfile.ZipFile(path) as original, zipfile.ZipFile(broken, 'w') as output:
            for entry in original.infolist():
                output.writestr(entry, b'<broken' if entry.filename == '[Content_Types].xml' else original.read(entry))
        before = broken.read_bytes()
        with self.assertRaises(ET.ParseError):
            remove_office_metadata(broken)
        self.assertEqual(broken.read_bytes(), before)
        self.assertFalse(list(self.folder.glob('.*_clean_*')))

    def test_backup_failure_does_not_modify_original(self):
        path = self.make_document('xlsx')
        before = path.read_bytes()
        with patch('cleaner.make_backup_copy', side_effect=PermissionError('denied')):
            with self.assertRaises(PermissionError):
                remove_office_metadata(path)
        self.assertEqual(path.read_bytes(), before)
        self.assertFalse(list(self.folder.glob('.*_clean_*')))

    def test_folder_scan_includes_all_families_ignores_office_lock_and_backup(self):
        for name in ('word.docx', 'slides.PPTX', 'sheet.XLSX', '~$lock.xlsx', 'old.xls', 'sheet.xlsx.bak'):
            (self.folder / name).touch()
        nested = self.folder / 'nested'
        nested.mkdir()
        (nested / 'inside.pptm').touch()
        self.assertEqual({p.name for p in iter_office_files(self.folder, False)}, {'word.docx', 'slides.PPTX', 'sheet.XLSX'})
        self.assertEqual(len(list(iter_office_files(self.folder, True))), 4)

    def test_nonstandard_properties_and_nested_relationships_are_removed(self):
        path = self.make_document('docx')
        special = self.folder / 'special.docx'
        with zipfile.ZipFile(path) as original, zipfile.ZipFile(special, 'w') as output:
            for entry in original.infolist():
                data = original.read(entry)
                if entry.filename == '_rels/.rels':
                    data = data.replace(b'docProps/core.xml', b'properties/author.xml')
                if entry.filename == '[Content_Types].xml':
                    data = data.replace(b'/docProps/core.xml', b'/properties/author.xml')
                if entry.filename == 'docProps/core.xml':
                    output.writestr('properties/author.xml', data)
                else:
                    output.writestr(entry, data)
            output.writestr('word/_rels/extra.xml.rels', b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="r1" Type="custom" Target="../properties/author.xml"/></Relationships>')
        remove_office_metadata(special, make_backup=False)
        with zipfile.ZipFile(special) as archive:
            self.assertNotIn('properties/author.xml', archive.namelist())
            self.assertEqual(len(ET.fromstring(archive.read('word/_rels/extra.xml.rels'))), 0)

    def test_settings_restore_desktop_and_migrate_old_version(self):
        desktop = self.folder / 'Рабочий стол'
        desktop.mkdir()
        settings_path = self.folder / 'config' / 'settings.json'
        save_settings(settings_path, {'last_directory': str(desktop), 'recursive': False, 'make_backup': False})
        loaded = load_settings(settings_path)
        self.assertEqual(restored_selection(loaded), (desktop.resolve(), 'folder'))
        save_settings(settings_path, dict(loaded, selected_path=str(desktop), selection_mode='folder'))
        self.assertEqual(restored_selection(load_settings(settings_path)), (desktop.resolve(), 'folder'))
        self.assertFalse(load_settings(settings_path)['make_backup'])
        self.assertTrue(desktop_directory().is_dir())

    def test_missing_selection_and_corrupt_settings(self):
        self.assertEqual(restored_selection({'selected_path': str(self.folder / 'gone')}), (None, ''))
        path = self.folder / 'settings.json'
        for text in ('{bad', '[]', 'null'):
            path.write_text(text)
            self.assertEqual(load_settings(path), {})

    def test_gui_restores_selected_desktop_and_cancel_keeps_it(self):
        from DocxMetaCleaner import DocxMetaCleanerApp
        settings_path = self.folder / 'settings.json'
        save_settings(settings_path, {'last_directory': str(self.folder), 'recursive': False, 'make_backup': False})
        app = DocxMetaCleanerApp(settings_path)
        app.withdraw()
        try:
            self.assertEqual(app.selected_path, self.folder.resolve())
            self.assertEqual(app.selection_mode.get(), 'folder')
            with patch('DocxMetaCleaner.filedialog.askdirectory', return_value=''):
                app.select_folder()
            self.assertEqual(app.selected_path, self.folder.resolve())
            app.select_desktop()
            expected = desktop_directory().resolve()
            self.assertEqual(app.selected_path, expected)
        finally:
            app.close()
        reopened = DocxMetaCleanerApp(settings_path)
        reopened.withdraw()
        try:
            self.assertEqual(reopened.selected_path, expected)
            self.assertEqual(reopened.path_var.get(), str(expected))
            self.assertEqual(reopened.selection_mode.get(), 'folder')
            self.assertFalse(reopened.backup_var.get())
        finally:
            reopened.close()


if __name__ == '__main__':
    unittest.main()
